"""Generate SRS.docx for TravelAI strictly from codebase inspection."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"C:\Users\pares\TravelAI\SRS.docx"

doc = Document()

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10)
for h in ['Heading 1','Heading 2','Heading 3']:
    s = doc.styles[h]
    s.font.name = 'Calibri'
    s.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

def add_title(t):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t)
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(0x1F,0x3A,0x5F)
    return p

def h1(t):
    return doc.add_heading(t, level=1)
def h2(t):
    return doc.add_heading(t, level=2)
def h3(t):
    return doc.add_heading(t, level=3)
def para(t, bold=False, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.bold = bold
    r.italic = italic
    return p
def bullets(items):
    for it in items:
        doc.add_paragraph(it, style='List Bullet')
def numbered(items):
    for it in items:
        doc.add_paragraph(it, style='List Number')

def shade_header(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), '1F3A5F')
    shd.set(qn('w:val'), 'clear')
    tcPr.append(shd)
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
            r.bold = True

def make_table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htxt in enumerate(headers):
        hdr[i].text = ''
        p = hdr[i].paragraphs[0]
        r = p.add_run(htxt)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
        shade_header(hdr[i])
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            r = p.add_run(str(val))
            r.font.size = Pt(8.5)
    doc.add_paragraph("")
    return t

# COVER
add_title("TravelAI")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Software Requirements Specification (SRS)\nVersion 1.0 — Grounded in Codebase Implementation")
r.font.size = Pt(12)
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run("Source: C:\\Users\\pares\\TravelAI  |  Date: 2026-09-19  |  Method: static inspection of Backend/, Frontend/, env examples, requirements, tests")
r2.font.size = Pt(8.5)
r2.italic = True
para("This SRS documents ONLY what exists in the code. No features were invented. Where a capability is absent in code (e.g., authentication, Docker, real payment/booking), it is explicitly marked as out-of-scope.", italic=True)

# TOC manual
h2("Table of Contents")
numbered([
    "1. Introduction",
    "2. Overall Description",
    "3. System Architecture",
    "4. Functional Requirements",
    "5. Non-Functional Requirements",
    "6. External Interface Requirements",
    "7. Database Requirements",
    "8. AI / ML Requirements",
    "9. API Requirements",
    "10. Deployment Requirements",
])

# 1
h1("1. Introduction")
h2("1.1 Purpose")
para("Specify the implemented behavior of TravelAI: a conversational AI trip planner with FastAPI backend, LangGraph multi-agent itinerary pipeline, RAG over ChromaDB, MongoDB persistence, and React + Vite + TypeScript frontend. Audience: developers, testers, maintainers onboarding to the actual codebase.")
h2("1.2 Scope")
bullets([
    "In scope (implemented): conversational chat with 4 intents (UPDATE_CONTEXT, TRAVEL_QA, GENERATE_ITINERARY, MODIFY_ITINERARY); 10-node LangGraph pipeline (parser, planner, 5 parallel specialists, budget, route, reviewer with rework loop); live data via SerpApi Google Flights (optional), Open-Meteo weather, Overpass OSM attractions/restaurants; RAG via ChromaDB + Ollama embeddings; MongoDB trips/conversations/messages; frontend Planner Studio, Itinerary Hub, Bookings Desk (client-side vouchers only), AI Places Memory (5 hardcoded Goa places).",
    "Out of scope (not in code): user authentication/login/JWT/OAuth, user accounts/roles, real flight/hotel payment or booking transactions, Docker/Kubernetes manifests, push notifications, admin console.",
])
h2("1.3 Definitions and Abbreviations")
make_table(["Term","Meaning (as used in code)"],[
    ["TravelState","LangGraph TypedDict in app/ai/orchestrator/state.py carrying user inputs, agent outputs, reviewer verdict, completed_steps, error_logs."],
    ["RAG","Retrieval-Augmented Generation via ChromaDB collection travel_knowledge + Ollama nomic-embed-text."],
    ["QA","TRAVEL_QA intent handled by qa_service.py bypassing the 10-node pipeline."],
    ["OSM / Overpass","OpenStreetMap Overpass API (overpass-api.de/api/interpreter) for attractions/restaurants."],
    ["SerpApi","SerpApi Google Flights engine (serpapi.com/search.json), requires SERPAPI_API_KEY; optional fallback chain."],
    ["TravelContext","Conversation-scoped slots in models/conversation.py (destination, origin, dates, budget, travelers, preferences, etc.)."],
])
h2("1.4 References")
bullets([
    "Backend/app/main.py (FastAPI title TravelAI API, CORS, routers).",
    "Backend/app/api/trips.py, app/api/chat.py; app/schemas/chat.py; app/models/trip.py, conversation.py, message.py.",
    "Backend/app/services/chat_service.py, conversation_service.py, slot_extraction_service.py, qa_service.py.",
    "Backend/app/ai/llm.py, orchestrator/{state,parser_agent,planner_agent,langgraph_graph}.py, agents/{flight,hotel,weather,attraction,restaurant,budget,route}_*.py, reviewer/reviewer_agent.py, rag/{chroma_client,embedder,retriever,ingest,download_india_guides}.py.",
    "Backend/requirements.txt (10 deps), .env.example (6 keys), test_graph.py, test_conversational_chat.py, chromadb_storage/, MVC ARCHITECTURE.png.",
    "Frontend/src/App.tsx, main.tsx, services/api.ts, hooks/useTravelApi.ts, types/index.ts, pages/{PlannerStudio,ItineraryHub,BookingsDesk,AiPlacesMemory}Page.tsx, components/{Sidebar,Header,ChatPanel,DashboardPanel,ApiConfigModal,ErrorBoundary}, tabs/{Route,Flights,Hotels,Weather,Budget}Tab.tsx, preferences/*Popover.tsx, package.json, vite.config.ts, .env.example.",
])
h2("1.5 Overview")
para("Sections 2–3 describe product and architecture. Section 4 lists functional requirements with IDs. Section 5 covers non-functionals observable from code (fallbacks, limits, UX). Sections 6–10 detail interfaces, data, AI, APIs, and deployment.")

# 2
h1("2. Overall Description")
h2("2.1 Product Perspective")
para("TravelAI is a client–server AI application. Frontend (React 19 + Vite 6 + Tailwind 4, package name react-example) talks REST/JSON to FastAPI backend (default http://localhost:8000). Backend orchestrates local Ollama LLM (qwen2.5:3b, temperature 0.0) and embeddings (nomic-embed-text), MongoDB (travel_ai), local ChromaDB (chromadb_storage), and public HTTP APIs. No auth layer; no gateway; CORS allows http://localhost:3000, http://127.0.0.1:3000, http://localhost:5173, http://127.0.0.1:5173. Frontend dev server runs vite --port=3000 --host=0.0.0.0 per package.json.")
h2("2.2 Product Functions (Summary)")
bullets([
    "Chat-driven slot capture, intent classification, and acknowledgement (chat_service + slot_extraction_service).",
    "One-shot or conversational full itinerary generation (LangGraph graph.invoke).",
    "Day-level itinerary modification of stored planner_draft via LLM (MODIFY_ITINERARY).",
    "Direct travel Q&A using live APIs or RAG without running full pipeline.",
    "Trip persistence + last-10 history; conversation/thread persistence with travel_context.",
    "Dashboard visualization: route, flights, hotels, weather, budget tabs; flight/hotel selection (client state); place append from memory.",
    "Mock/preset fallback (INITIAL_*, PRESET_TRIPS: Goa, Tokyo, Rajasthan) + localStorage caches when backend unreachable; API health indicator + configurable base URL.",
])
h2("2.3 User Characteristics")
bullets([
    "End user: traveler using browser chat + preference bar (Where/When/Budget/Travelers/Interests) to build/modify trips; needs no login.",
    "Maintainer: runs Ollama, MongoDB, Python backend, Node frontend; manages .env keys and knowledge/*.txt ingestion.",
])
h2("2.4 Constraints")
bullets([
    "LLM/embedding availability: Ollama at OLLAMA_BASE_URL must serve qwen2.5:3b and nomic-embed-text; all agents block on llm.invoke.",
    "MongoDB at MONGODB_URI/DB_NAME; Chroma persist dir Backend/chromadb_storage must be writable.",
    "SerpApi key optional — flight agent degrades to verified schedules then RAG/LLM then hardcoded fallback if missing/failing.",
    "City coverage for live geo: hardcoded coordinates for ~16 Indian cities + airport map for ~20 names; unknown cities fall back to defaults (Central India coords, GOI/HYD codes).",
    "Reviewer loop capped at 3 planner revisits to avoid infinite loops.",
])
h2("2.5 Assumptions and Dependencies")
bullets([
    "Python deps: fastapi, uvicorn, langchain, langgraph, langchain-ollama, langchain-community, pymongo, chromadb, httpx, python-dotenv.",
    "Node deps include react, react-dom, vite, tailwind, lucide-react, motion, express, dotenv, @google/genai (declared; no direct AI call observed in inspected frontend services — chat goes via FastAPI).",
    "External HTTP reachable: api.open-meteo.com, overpass-api.de, serpapi.com (if key), en.wikivoyage.org (ingest script only), google.com/travel/flights (constructed links only).",
    "Frontend env VITE_FASTAPI_URL (or VITE_API_URL) defaults to http://localhost:8000; runtime override via localStorage TRAVELAI_FASTAPI_URL.",
])
h2("2.6 Explicit Non-Goals (Verified Absent)")
bullets([
    "No authentication/authorization code (no login, JWT, sessions, password hashing) in Backend or Frontend.",
    "No Dockerfile, docker-compose, CI YAML, Procfile, or shell deploy scripts found via glob.",
    "No real booking/payment: BookingsDeskPage banner states 'No external booking API required'; vouchers/PNRs are display-only; flight links are Google search URLs.",
])

# 3
h1("3. System Architecture")
h2("3.1 Logical Architecture")
para("Three tiers + sidecar AI/data services:")
bullets([
    "Presentation: React SPA (App.tsx switch on activePage: planner-studio | itinerary-hub | bookings-desk | ai-places-memory; Sidebar + Header + ApiConfigModal + ErrorBoundary). State via useTravelApi hook; pure mappers in services/api.ts (fetchFlightOptions, fetchHotelOptions, fetchWeatherCondition, fetchBudgetOverview, extractTripSummary, extractItineraryDays).",
    "Application/API: FastAPI app travel_ai in app/main.py; routers trips (/api/trips) and chat (/api/chat); services layer (chat, conversation, slot_extraction, qa).",
    "AI orchestration: LangGraph StateGraph (TravelState) with 10 nodes; shared llm instance (ChatOllama temp 0.0).",
    "Data: MongoDB (pymongo MongoClient, db=travel_ai) + ChromaDB (LangChain Chroma wrapper, collection travel_knowledge, persist chromadb_storage).",
    "Integrations: Ollama LLM/embeddings, SerpApi, Open-Meteo, Overpass, Wikivoyage (offline ingest). Repo also contains MVC ARCHITECTURE.png (diagram asset, not parsed).",
])
h2("3.2 Backend Module Map")
make_table(["Layer","Files","Responsibility"],[
    ["Entry","app/main.py","FastAPI init, CORS, include routers, GET / -> {status,message}."],
    ["API","app/api/trips.py, chat.py","/plan, /history, /message, /conversations*, validation via Pydantic."],
    ["Schemas/Models","app/schemas/chat.py; app/models/*.py","Request/response shapes; TripModel, ConversationModel+TravelContext, MessageModel."],
    ["Services","app/services/*.py","Conversation CRUD+indexes; slot LLM+regex; chat intent dispatch; QA synthesis."],
    ["AI core","app/ai/llm.py; orchestrator/*; agents/*; reviewer/*","LLM singleton; parser/planner; 7 specialists; reviewer; graph wiring."],
    ["RAG","app/ai/rag/*.py","embeddings, Chroma client, retriever+seed, ingest splitter, Wikivoyage downloader."],
    ["DB","app/database/mongodb.py","MongoClient(MONGODB_URI), db handle, connect log."],
])
h2("3.3 LangGraph Workflow (as coded)")
para("Nodes: parser → planner → {flights, hotels, weather, attractions, restaurants} (5 parallel) → budget → route → reviewer → END or → planner (conditional). Entry: parser. Reviewer should_continue(): if is_approved true → approved/END; else if planner count ≥3 → force END; else rejected → planner.")
bullets([
    "parser_node: skips LLM if destination+days+budget present; else Qwen JSON extract {destination,days,budget,preferences} + normalize_currency_and_numbers; defaults days 3, budget 500.0, destination Unknown on error.",
    "planner_node: Qwen day-by-day draft from {destination,days,budget,preferences}.",
    "Specialists run in parallel reading planner_draft + RAG context; each returns its slice of TravelState.",
    "budget_node: Qwen audit → {flight_cost,hotel_cost,gastronomy_cost,activities_cost,unallocated_buffer,total_trip_cost,remaining_budget,is_under_budget} + travelers/currency/totalBudget; deterministic percentage fallback on parse error.",
    "route_node: Qwen day_by_day_route {day,stops,travel_tips}; fallback builds per-day stops from first hotel/attraction/restaurant.",
    "reviewer_node: Qwen {is_approved,reviewer_feedback} checking budget, destination match, route sense; parse failure → is_approved False.",
])
h2("3.4 Frontend Architecture")
bullets([
    "AppContent holds all state via useTravelApi; auto-dismiss error toast 5s; handleNewTrip → handleNewChat + focus chatPromptInput.",
    "PlannerStudioPage: left ChatPanel (messages, promptInput, travelContext, intent, preference bar) + right DashboardPanel (5 tabs). Empty state if !trip.",
    "Dashboard tabs: RouteTab, FlightsTab, HotelsTab, WeatherTab, BudgetTab fed by trip/days/flights/hotels/weather/budget props.",
    "ItineraryHubPage: search by title/destination/preferences, filter all/active/archived, load into studio, delete (localStorage TRAVELAI_DELETED_TRIP_IDS).",
    "BookingsDeskPage: derives selectedFlight/bookedHotel from isSelected/isBooked; static chauffeur card; copy/share/print via clipboard/window.print.",
    "AiPlacesMemoryPage: 5 static places (Cavatina, Thalassa, Bom Jesus, Fontainhas, Morjim); tag filter all/dining/beach/heritage/saved + search; save toggle + Append to Route (adds ActivitySegment to day 2 or day 1).",
    "Sidebar: nav + Recent Chats (/api/chat) + Saved Trips (/api/trips/history) + API status footer + New Bespoke Trip CTA; Header: search + trip presets + mobile toggle.",
])
h2("3.5 Data Flow (Implemented)")
numbered([
    "User types in ChatPanel → useTravelApi.handleSendMessage → local checkTripPlanningConstraints guidance or POST /api/chat/message {message, conversation_id}.",
    "Backend chat_service: get_or_create_conversation → add user msg → analyze_user_message (LLM ANALYZER_PROMPT + regex fallback + modify override) → branch by intent.",
    "GENERATE_ITINERARY: build_travel_state_from_context (defaults Goa/Hyderabad/3d/budget by style) → graph.invoke → save trips doc → update conversation current_trip_id → return {conversation_id,message,intent,travel_context,data:{type:trip_id,itinerary}}.",
    "Frontend syncTripState maps raw trip via extract*/fetch* helpers; tabs render; history/sidebar refresh.",
])

# 4
h1("4. Functional Requirements")
para("IDs FR-1..FR-14 are verifiable against code paths named in brackets.")
make_table(["ID","Requirement","Source / Acceptance"],[
    ["FR-1","Conversational chat accepts free text and returns assistant reply + intent + travel_context. Empty message rejected 400.","chat.py send_chat_message; chat_service.process_chat_message; test_conversational_chat 8 turns"],
    ["FR-2","Classify intent into UPDATE_CONTEXT / TRAVEL_QA / GENERATE_ITINERARY / MODIFY_ITINERARY with slots, qa_category, confidence.","slot_extraction ANALYZER_PROMPT; deterministic modify regex when has_current_trip"],
    ["FR-3","UPDATE_CONTEXT merges slots non-destructively (scalars overwrite, lists append unique) and LLM-acknowledges (1–2 sentences).","merge_travel_context; ACKNOWLEDGEMENT_PROMPT; title auto Trip to {Destination}"],
    ["FR-4","GENERATE_ITINERARY requires destination+days else asks for missing; else runs full LangGraph and persists trip.","chat_service missing check; build_travel_state_from_context defaults; graph.invoke"],
    ["FR-5","MODIFY_ITINERARY requires existing current_trip_id; LLM rewrites planner_draft only for requested change; updates trips doc.","MODIFIER_PROMPT; db.trips.update_one planner_draft"],
    ["FR-6","TRAVEL_QA answers via live weather/OSM or Chroma RAG + LLM synthesis without pipeline.","qa_service.handle_travel_qa categories WEATHER/ATTRACTION/RESTAURANT/GENERAL"],
    ["FR-7","One-shot plan API POST /api/trips/plan accepts {user_message|destination,days,budget,preferences} and returns full trip doc with id.","trips.py plan_trip; TripModel fields"],
    ["FR-8","History APIs: GET /api/trips/history (latest 10), GET/DELETE /api/chat/conversations[/{id}], GET / health.","trips.py get_trip_history; chat.py list/detail/delete; main.py read_root"],
    ["FR-9","Dashboard renders route/flights/hotels/weather/budget from active trip; flight/hotel single-select; deterministic mappers with fallbacks.","DashboardPanel + tabs; api.ts fetch*/extract* + INITIAL_* fallbacks"],
    ["FR-10","Preference bar edits travelContext (Where/When/Budget/Travelers/Interests) client-side and enriches prompt via [Context:] tags + parsers.","TravelPreferenceBar + Where/When/Budget/Travelers/InterestsPopover; parseExplicitChatbotSlots, buildChatPromptWithFallbackContext"],
    ["FR-11","Itinerary Hub lists merged backend history + presets, search/filter, open in studio, delete.","ItineraryHubPage; PRESET_TRIPS Goa/Tokyo/Rajasthan; fetchTripHistory"],
    ["FR-12","Bookings Desk shows derived flight/hotel/transfer vouchers, share/copy/print; no server booking call.","BookingsDeskPage; banner No external booking API required"],
    ["FR-13","AI Places Memory lists 5 curated places, save/unsave, search/filter, append to active route.","AiPlacesMemoryPage; handleAppendPlaceToRoute"],
    ["FR-14","API connectivity: health GET / (5s timeout), status listener live/mock-fallback, configurable URL modal, 25s re-ping, localStorage caches.","api.ts checkFastApiHealth/apiFetch; ApiConfigModal; useTravelApi loadData"],
])
h3("4.1 Conversation Persistence Details")
bullets([
    "IDs: conv_{10hex}, msg_{10hex}; messages store {message_id,conversation_id,role,user|assistant|system,content,intent,structured_data,created_at}.",
    "list_conversations: status=active sorted updated_at desc limit 20; get_messages sorted created_at asc limit 50; delete removes conversation + messages.",
    "Frontend caches TRAVELAI_CONVERSATION_ID, TRAVELAI_ACTIVE_MESSAGES, TRAVELAI_TRAVEL_CONTEXT, TRAVELAI_CURRENT_INTENT, TRAVELAI_TRIP_CONTEXTS, deleted IDs.",
])

# 5
h1("5. Non-Functional Requirements")
make_table(["Category","Requirement (observable)"],[
    ["Performance","Health check timeout 5s; chat/plan fetches have no client timeout (await completion); parallel specialist nodes; Overpass/SerpApi/Open-Meteo httpx timeouts 10–15s; ingest batches 100 docs."],
    ["Reliability","Per-agent fallback chains (live → verified/RAG → hardcoded); parser/budget/route/reviewer JSON parse guards with defaults; DB save failure returns temporary_id yet returns data; ErrorBoundary per page + toast."],
    ["Correctness guard","Reviewer rework loop max 3 planner passes; budget recomputed deterministically (buffer = budget − sum, is_under_budget flag)."],
    ["Usability","Responsive layout (mobile sidebar/backdrop, collapse), loading skeletons, empty states, toasts, 5s auto-dismiss error banner, aria labels on key buttons."],
    ["Determinism","LLM temperature 0.0; regex normalizers for INR/USD/lakh/k, dates, travelers, [Context:] tags ensure stable slots when LLM misses."],
    ["Maintainability","Typed Pydantic models + TS interfaces; modular agents/services; test scripts test_graph.py and test_conversational_chat.py with asserts."],
    ["Portability","CORS for Vite 5173 + CRA 3000; backend/frontend .env examples; no OS-specific calls except paths."],
    ["Security/Privacy (as-is)","No auth — single-tenant local use; secrets via .env (gitignored); no password/PII handling beyond travel prefs stored in Mongo/localStorage."],
])

# 6
h1("6. External Interface Requirements")
h2("6.1 User Interfaces")
bullets([
    "Planner Studio (PlannerStudioPage): ChatPanel left, DashboardPanel right with tabs route-tab|flights-tab|hotels-tab|weather-tab|budget-tab.",
    "Itinerary Hub: cards with image, status, destination, duration, budget, travelers, style/mode chips; search + all/active/archived filter.",
    "Bookings Desk: All/Flights/Hotels/Transfers filter; flight PNR #6E9482 display, hotel Res #TAJ8830 display, chauffeur card (static).",
    "AI Places Memory: place cards with rating/location/price + Save heart + Append to Route.",
    "Sidebar/Header/ApiConfigModal: navigation, recent chats, saved trips, API status (live/mock-fallback), base-URL editor.",
])
h2("6.2 Hardware Interfaces")
para("None beyond standard browser client and server host running Ollama/MongoDB. No sensor/device integration in code.")
h2("6.3 Software Interfaces")
make_table(["Peer","Interface (coded)"],[
    ["Ollama LLM","ChatOllama base_url=OLLAMA_BASE_URL model=qwen2.5:3b temp 0.0 (llm.py)."],
    ["Ollama Embeddings","OllamaEmbeddings model nomic-embed-text (embedder.py)."],
    ["MongoDB","pymongo MongoClient(MONGODB_URI) db travel_ai (mongodb.py)."],
    ["ChromaDB","LangChain Chroma collection travel_knowledge persist_directory chromadb_storage (chroma_client.py)."],
    ["SerpApi","GET https://serpapi.com/search.json engine=google_flights departure_id/arrival_id/currency (+outbound_date) via httpx."],
    ["Open-Meteo","GET https://api.open-meteo.com/v1/forecast current+daily forecast_days min(days,7)."],
    ["Overpass OSM","POST https://overpass-api.de/api/interpreter tourism/historic (attractions) and amenity restaurant/cafe (restaurants)."],
    ["Wikivoyage","GET https://en.wikivoyage.org/w/api.php action=query prop=extracts (download script only)."],
    ["Google Flights links","Frontend constructs https://www.google.com/travel/flights?q=... and search URLs (no API key)."],
])
h2("6.4 Communications Interfaces")
bullets([
    "REST/JSON over HTTP; Content-Type/Accept application/json; apiFetch throws Error(detail) on !response.ok parsing {detail}.",
    "CORS allow_origins 4 localhost entries, allow_credentials true, allow_methods/headers *.",
    "Health: GET / → {status:online,message}. No /api/health endpoint in code.",
])

# 7
h1("7. Database Requirements")
h2("7.1 MongoDB (travel_ai)")
para("Client: MongoClient(MONGODB_URI default mongodb://localhost:27017); DB_NAME default travel_ai. Collections observed: trips, conversations, messages (via db.trips, db[conversations], db[messages]).")
make_table(["Collection","Key Fields (from models/code)"],[
    ["trips","user_message, destination (default Unknown), days (default 3), budget (default 500.0), preferences, planner_draft, flights[], hotels[], weather{}, attractions[], restaurants[], route_details{}, budget_breakdown{}, origin/travelers/currency/travel_style/budget_mode/start_date/end_date/travelContext (chat path), created_at, _id→id string."],
    ["conversations","conversation_id (unique idx), title (auto Trip to X), status=active, travel_context{TravelContext}, current_trip_id, created_at, updated_at (desc idx)."],
    ["messages","message_id, conversation_id, role, content, intent, structured_data{trip_id|slots}, created_at; idx (conversation_id,created_at)."],
])
bullets([
    "Indexes (conversation_service.setup_indexes): conversations(conversation_id unique), conversations(updated_at -1), messages(conversation_id 1, created_at 1).",
    "Ops: trips.insert_one + find().sort(created_at -1).limit(10); conversations insert/find/update/delete; messages insert/find/delete_many.",
    "Failure semantics: plan/chat catch DB errors, log warning, return temporary_id/temp_trip_id with full payload (no hard fail).",
])
h2("7.2 ChromaDB Vector Store")
bullets([
    "Persist dir: Backend/chromadb_storage (chroma.sqlite3 + UUID subdir observed); collection travel_knowledge via LangChain Chroma.",
    "Seed if count==0: 5 Goa facts (retriever.MOCK_TRAVEL_DATA: flights GOI/GOX, North/South hotels, dining $10–15, scooter $5–8/day, monsoon Jun–Sep ideal Oct–Feb).",
    "Retrieval: similarity_search(query,k) k=2 specialists/QA, k=3 general QA; returns newline-joined page_content; empty string on error.",
    "Ingest: knowledge/*.txt via DirectoryLoader+TextLoader utf-8 → RecursiveCharacterTextSplitter chunk 500 overlap 50 → add_documents batches 100; creates knowledge/ if missing.",
    "Knowledge bootstrap: download_india_guides.py fetches 15 cities (Agra, Jaipur, Delhi, Mumbai, Bengaluru, Kerala, Varanasi, Udaipur, Leh, Manali, Shimla, Amritsar, Ooty, Darjeeling, Goa) to knowledge/{city}_guide.txt.",
])

# 8
h1("8. AI / ML Requirements")
h2("8.1 LLM & Embeddings")
make_table(["Item","Coded Value"],[
    ["Chat model","OLLAMA_MODEL default qwen2.5:3b via ChatOllama, temperature 0.0 (app/ai/llm.py)."],
    ["Embeddings","nomic-embed-text via OllamaEmbeddings at OLLAMA_BASE_URL."],
    ["Prompt framework","langchain_core.prompts.ChatPromptTemplate system+human; raw JSON-only responses; manual ``` strip + json.loads."],
])
h2("8.2 Prompts Inventory")
bullets([
    "PARSER_PROMPT (parser_agent): extract {destination,days,',budget,preferences} defaults days3 budget500.",
    "PLANNER_PROMPT: day-by-day draft from {destination,days,budget,preferences}.",
    "ANALYZER_PROMPT (slot_extraction): intent+slots+qa_category+confidence+clarification; slot rules for budget_mode/travel_style/transport/accommodation/preferences/constraints.",
    "ACKNOWLEDGEMENT_PROMPT + MODIFIER_PROMPT (chat_service): warm ack; draft-only modification.",
    "QA_SYNTHESIS_PROMPT: answer from {context}.",
    "FLIGHT/HOTEL/WEATHER/ATTRACTION/RESTAURANT/BUDGET/ROUTE/REVIEWER prompts per agent with strict JSON keys (see agent table).",
])
h2("8.3 Agents Table")
make_table(["Agent","Inputs → Outputs","External/RAG + Fallback"],[
    ["parser","user_message → destination,days,budget,preferences","Qwen + normalize_currency_and_numbers; skip if structured"],
    ["planner","destination,days,budget,preferences → planner_draft","Qwen only"],
    ["flight","origin,destination,planner_draft,budget,currency,start_date → flights[{airline,flight_no,price,departure,arrival,duration,origin,destination}]","SerpApi → VERIFIED_SCHEDULES (8 routes HYD-COK/DEL/GOI/BOM/BLR etc.) → RAG+Qwen → 3 hardcoded; INR 3500–7500 realistic"],
    ["hotel","destination,days,planner_draft,budget,currency → hotels[{name,rating,price_per_night,total_cost,address,amenities}]","RAG+Qwen → 2 hardcoded; INR 4500–14000/night"],
    ["weather","destination,days → weather{average_temp,condition,clothing_recommendation,summary}","Open-Meteo + WMO codes + clothing logic → RAG+Qwen → Pleasant default"],
    ["attraction","destination,planner_draft,preferences → attractions[{name,description,entrance_fee,rating,recommended_time_spent}]","Overpass tourism/historic 8km/20 → RAG+Qwen → 1 fallback"],
    ["restaurant","destination,planner_draft,preferences → restaurants[{name,cuisine,price_range,rating,address,popular_dish}]","Overpass amenity 5km/15 max4 → RAG+Qwen → 1 fallback"],
    ["budget","budget,currency,travelers,budget_mode,travel_style,days,flights,hotels,attractions → breakdown{8 keys}","Qwen + deterministic recompute; fallback 25/40/18/10% split"],
    ["route","hotels,attractions,restaurants → route_details{day_by_day_route[{day,stops,travel_tips}]}","Qwen; fallback per-day firsts"],
    ["reviewer","all outputs + constraints → {is_approved,reviewer_feedback}","Qwen; parse fail → rejected; loop ≤3"],
])
h2("8.4 Slot & Geo Knowledge")
bullets([
    "Currency regex: lakh×100000, k×1000, ₹/INR/RS, $/USD; [Context: origin|destination|duration|dates|budget|travelers|interests] tags; date-range + pax regex; budget_mode/travel_style inference (cheap→ECONOMY/BUDGET, rich/no-limit→LUXURY).",
    "AIRPORT_CODES (~20 entries) + CITY_COORDINATES (16 cities) duplicated in flight/weather/attraction/restaurant agents; frontend resolveAirportCode mirrors + Google/HYD-COK mismatch guard.",
    "TravelContext defaults in chat bridge: origin Hyderabad, dest Goa, days 3, travelers 2, currency INR, budget by style (LUXURY 150000/2000, BUDGET 20000/300, else 50000/750).",
])

# 9
h1("9. API Requirements")
h2("9.1 Endpoints (as routed)")
make_table(["Method + Path","Request (Pydantic)","Response / Behavior"],[
    ["GET /","—","{status:online, message:TravelAI Backend is running successfully!} used as health."],
    ["POST /api/trips/plan","PlanRequest{user_message?,destination?,days?,budget?,preferences?}","Full trip doc + id (or temporary_id on DB fail); 500 on graph failure."],
    ["GET /api/trips/history","—","Array latest 10 trips, _id→id; 500 on DB fail."],
    ["POST /api/chat/message","ChatMessageRequest{conversation_id?,message}","ChatMessageResponse{conversation_id,message{role,content},intent,travel_context,data?{type,trip_id,itinerary}}; 400 if empty; 500 on error."],
    ["GET /api/chat/conversations","—","List ConversationSummaryResponse{conversation_id,title,status,updated_at,current_trip_id} limit 20."],
    ["GET /api/chat/conversations/{id}","path id","ConversationDetailResponse{...travel_context,messages[],created_at,updated_at}; 404 if missing."],
    ["DELETE /api/chat/conversations/{id}","path id","{status:deleted,conversation_id}; 404 if missing."],
])
h2("9.2 Schemas & Validation")
bullets([
    "ChatMessageRequest requires non-empty trimmed message; PlanRequest all-optional (parser fills defaults).",
    "TravelContext fields: destination, origin, start_date, end_date, days, budget, currency default INR, budget_mode FIXED/FLEXIBLE/ECONOMY/PREMIUM/LUXURY/NO_LIMIT, travel_style BUDGET/STANDARD/PREMIUM/LUXURY, travelers default 1, preferences/constraints/accommodation_preferences/transport_preferences lists, current_trip_id.",
    "TripModel vs chat trip_data: chat path adds origin, travelers/travelersCount, currency, travel_style, budget_mode, start/end_date, travelContext.",
    "Frontend contract: POST body EXACTLY {message, conversation_id} (useTravelApi comment); apiFetch sets JSON headers, parses {detail} on error, flips status to mock-fallback on network fail.",
])
h2("9.3 Frontend Service Functions")
bullets([
    "apiFetch<T>(endpoint, {timeoutMs?}), get/setApiBaseUrl, checkFastApiHealth, fetchTripHistory, sendChatMessage, fetchConversationsList/Detail, deleteConversation.",
    "Pure mappers: getTripStartDate, resolveAirportCode, getRouteScheduleData, computeTripDates, fetchFlightOptions/HotelOptions/WeatherCondition/BudgetOverview, extractTripSummary/ItineraryDays, parseMetadataFromUserMessage, parseExplicitChatbotSlots, mergeTravelContexts, checkTripPlanningConstraints, buildChatPromptWithFallbackContext.",
])

# 10
h1("10. Deployment Requirements")
h2("10.1 Prerequisites")
make_table(["Component","Version / Notes (from code)"],[
    ["Python","3.13 bytecode observed (.pyc cpython-313); install requirements.txt 10 packages."],
    ["Node.js","For npm install + npm run dev/build/preview/lint; Vite 6, React 19, TS ~5.8."],
    ["Ollama","Serve qwen2.5:3b + nomic-embed-text at OLLAMA_BASE_URL."],
    ["MongoDB","At MONGODB_URI, DB travel_ai; conversational + trips persistence."],
    ["Browser","Modern browser for SPA + localStorage."],
])
h2("10.2 Environment Variables")
make_table(["File","Key","Default / Purpose"],[
    ["Backend/.env","OLLAMA_BASE_URL","http://localhost:11434"],
    ["Backend/.env","OLLAMA_MODEL","qwen2.5:3b"],
    ["Backend/.env","MONGODB_URI","mongodb://localhost:27017"],
    ["Backend/.env","MONGODB_DB_NAME","travel_ai"],
    ["Backend/.env","SERPAPI_API_KEY","(optional) enables live flights; absent → fallback chain"],
    ["Backend/.env","HOTEL_API_KEY","Declared in example; no usage observed in inspected agents"],
    ["Frontend/.env","VITE_FASTAPI_URL / VITE_API_URL","http://localhost:8000; override via localStorage TRAVELAI_FASTAPI_URL"],
    ["Frontend/.env.example","GEMINI_API_KEY / APP_URL","Documented for AI Studio; no backend route or frontend service call observed using them"],
])
h2("10.3 Run & Build (inferred from scripts, no Docker)")
numbered([
    "Backend: set .env from .env.example → start Ollama (qwen2.5:3b, nomic-embed-text) → start MongoDB → run FastAPI (uvicorn, e.g., uvicorn app.main:travel_ai --reload --port 8000; module var is travel_ai not app). Verify GET /.",
    "RAG (optional): python app/ai/rag/download_india_guides.py → python app/ai/rag/ingest.py (chunks 500/50, batches 100) → chromadb_storage populated.",
    "Frontend: npm install → set VITE_FASTAPI_URL → npm run dev (vite --port=3000 --host=0.0.0.0) → open Planner Studio; health re-pings every 25s; ApiConfigModal can switch URL at runtime.",
    "Tests: python test_graph.py (pipeline e2e) and python test_conversational_chat.py (8-turn asserts UPDATE_CONTEXT/TRAVEL_QA/GENERATE/MODIFY).",
    "Build: npm run build → dist/ (dist/index.html + assets observed); backend has no build step.",
])
h2("10.4 Operational Notes")
bullets([
    "Ports: backend 8000, frontend dev 3000, Vite default 5173 also CORS-allowed; Ollama 11434.",
    "No auth, rate-limit, logging framework, or container orchestration in code; logs are print() statements with --- AGENT: prefixes.",
    "Git ignores: .venv/venv, .env + Backend/.env, __pycache__, node_modules/, dist/, .vite/.next/.cache, IDE files.",
])

# Sign-off
h1("Appendix — Traceability Note")
para("Every statement above maps to a file/line inspected on 2026-09-19. If code changes, regenerate this SRS. No hidden features: authentication, payments, Docker, notifications, and admin roles were searched (grep auth/login/jwt/token + glob Dockerfile/compose/yaml) and not found.")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("— End of SRS v1.0 —")
r.bold = True
r.font.size = Pt(10)

doc.save(OUT)
print(f"Saved {OUT}")
